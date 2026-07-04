"""Rockland green-sheet demo harness -- seed an ISOLATED demo DB, then serve it.

Why this exists (owner ask, 2026-07-04): a live demo/trial must never write mock
dollars into the canonical `project_db.sqlite`. This script copies the real DB to
`project_db.demo.sqlite`, seeds the pilot walkthrough there, and serves the UI
against that copy. Reset = just re-run seed (fresh copy every time).

What the seed does -- deliberately through the FRONT DOOR, exercising the whole
spine end-to-end rather than inserting rows by hand:

  1. BudgetSnapshot "v1 (pilot mock)" from the owner-approved mock
     `2026001_BUDGET_v1.xlsx` (12 division targets).
  2. Parses the two mock QUOTE workbooks with the real XlsxParser
     (DocumentParse + EvidenceSpan).
  3. Resolves each filename with `resolve_quote_document` (project/package/
     vendor from the filename alone -- Phase 5 item #3).
  4. Ingests via `ingest_subcontractor_quote` (cost rows, SOW_Item_Ref links).
  5. Awards the PO for the *selected* plumbing quote (`award_purchase_order`)
     -> cost flips quoted -> committed + ContractObligation emitted.

PROVENANCE: every dollar seeded here is a MOCK example from the approved
template drive (docs/templates/mock_drive). Nothing here is real financial
data; the demo DB is disposable and gitignored (*.sqlite).

Usage (from anywhere):
    python project-db/scripts/demo_rockland.py            # seed (fresh copy)
    python project-db/scripts/demo_rockland.py serve      # seed if missing + serve
    python project-db/scripts/demo_rockland.py serve --port 8123 --reseed
"""

from __future__ import annotations

import argparse
import os
import shutil
import sys
from pathlib import Path

_HERE = Path(__file__).resolve().parent
PROJECT_DB_DIR = _HERE.parent  # .../project-db
REPO_ROOT = PROJECT_DB_DIR.parent
REAL_DB = PROJECT_DB_DIR / "project_db.sqlite"
DEMO_DB = PROJECT_DB_DIR / "project_db.demo.sqlite"
MOCK_DRIVE = REPO_ROOT / "docs" / "templates" / "mock_drive"

# The demo engine must point at the demo DB BEFORE any project_db.db import.
os.environ["PROJECT_DB_URL"] = f"sqlite:///{DEMO_DB.as_posix()}"

sys.path.insert(0, str(PROJECT_DB_DIR / "src"))

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _mock(pattern: str) -> Path:
    hit = next(MOCK_DRIVE.glob(pattern), None)
    if hit is None:
        raise SystemExit(f"mock file not found: {pattern} under {MOCK_DRIVE}")
    return hit


def seed() -> None:
    from project_db.cli import force_utf8_output

    force_utf8_output()

    if not REAL_DB.exists():
        raise SystemExit(f"real DB not found at {REAL_DB} -- nothing to copy")
    if DEMO_DB.exists():
        DEMO_DB.unlink()
    shutil.copyfile(REAL_DB, DEMO_DB)
    print(f"OK: copied real DB -> {DEMO_DB.name} (isolated; canonical DB untouched)")

    import openpyxl

    from project_db.ai.financial_divisions import canonical_division_code
    from project_db.ai.green_sheet import report_green_sheet
    from project_db.ai.purchase_order_award import award_purchase_order
    from project_db.ai.quote_document_resolver import resolve_quote_document
    from project_db.ai.subcontractor_quote_ingest import ingest_subcontractor_quote
    from project_db.db import Base, ensure_sqlite_schema, get_engine, session_scope
    from project_db.db.models import (
        BudgetSnapshot,
        BudgetSnapshotLine,
        Document,
        Organization,
        Project,
        SubcontractorQuote,
        Vendor,
    )
    from project_db.parsing.service import parse_document_content

    engine = get_engine()
    assert "demo" in str(engine.url), f"refusing: engine is not the demo DB ({engine.url})"
    Base.metadata.create_all(engine)
    ensure_sqlite_schema(engine)

    budget_xlsx = _mock("*/budget/2026001_BUDGET_v1.xlsx")
    quote_files = [
        _mock("*/quotes/2026001_QUOTE_22-Plumbing_PlombertInc_selected.xlsx"),
        _mock("*/quotes/2026001_QUOTE_09-Finishes_ABCTile_pending.xlsx"),
    ]

    with session_scope() as s:
        project = s.query(Project).filter(Project.code == "2026001").one()
        org = s.query(Organization).first()

        # Vendors the resolver will match by VendorSlug (never guessed).
        for name in ("Plombert Inc.", "ABC Tile"):
            if s.query(Vendor).filter(Vendor.name == name).one_or_none() is None:
                s.add(Vendor(name=name, organization_id=org.canonical_id))
        s.flush()

        # 1. Budget baseline from the approved mock BUDGET workbook.
        snap = BudgetSnapshot(
            project_id=project.canonical_id,
            label="v1 (pilot mock)",
            source_meta_json='{"provenance": "mock 2026001_BUDGET_v1.xlsx -- demo only"}',
        )
        s.add(snap)
        s.flush()
        ws = openpyxl.load_workbook(budget_xlsx, data_only=True)["Budget_Lines"]
        n_budget = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            line_id, div_code, trade, amount, markup, _client = row
            if line_id == "TOTAL" or div_code is None:
                continue
            s.add(
                BudgetSnapshotLine(
                    snapshot_id=snap.canonical_id,
                    project_id=project.canonical_id,
                    division_code=canonical_division_code(str(div_code)),
                    division_name=trade,
                    budget_amount=amount,
                    line_markup_factor=markup,
                )
            )
            n_budget += 1
        print(f"OK: BudgetSnapshot 'v1 (pilot mock)' seeded ({n_budget} division targets)")

        # 2-4. Each mock quote: parse -> resolve from filename -> ingest.
        for qf in quote_files:
            doc = Document(
                name=qf.name,
                url=f"file:///{qf.name}",
                project_id=project.canonical_id,
                mime_type=XLSX_MIME,
            )
            s.add(doc)
            s.flush()
            parse = parse_document_content(
                s, document=doc, content=qf.read_bytes(), filename=qf.name
            )
            res = resolve_quote_document(s, qf.name)
            if not res.fully_resolved:
                raise SystemExit(f"resolver did not fully resolve {qf.name}: {res.warnings}")
            ing = ingest_subcontractor_quote(
                s,
                doc,
                project_id=res.project_id,
                package_id=res.package_id,
                vendor_id=res.vendor_id,
                division_code=res.parsed.division_code,
            )
            s.flush()
            print(
                f"OK: {qf.name}: parse={parse.status}, resolved via filename "
                f"({res.package_method}/{res.vendor_method}), "
                f"{ing.rows_written} cost rows, status={ing.status}, "
                f"reconcile_ok={ing.reconcile_ok}"
            )

        # 5. Award the PO for the selected plumbing quote.
        selected = (
            s.query(SubcontractorQuote)
            .filter(
                SubcontractorQuote.project_id == project.canonical_id,
                SubcontractorQuote.status == "selected",
            )
            .one()
        )
        po = award_purchase_order(s, selected)
        print(f"OK: awarded PO {po.po_number} ({po.lines_committed} cost rows -> committed)")

        gs = report_green_sheet(s, "2026001")
        print(
            f"\nGreen sheet: budget={gs['total_budget']}, quoted={gs['total_quoted']}, "
            f"pending_bids={gs['total_pending_bids']}, committed={gs['total_committed']}, "
            f"variance={gs['total_variance']}"
        )
        print(f"\nDemo DB ready: {DEMO_DB}")
        print("Serve it:   python scripts/demo_rockland.py serve")
        print(f"Then open:  http://127.0.0.1:8123/projects/{project.canonical_id}/green-sheet")


def serve(port: int, reseed: bool) -> None:
    if reseed or not DEMO_DB.exists():
        seed()
    try:
        import uvicorn
    except ImportError:
        raise SystemExit('uvicorn not installed. Install the UI extra: pip install -e ".[ui]"')

    from project_db.web.app import create_app

    # No background refresh on purpose: the demo DB must never delta-sync.
    print(f"\nDEMO UI (isolated {DEMO_DB.name}) on http://127.0.0.1:{port}  -- Ctrl-C to stop")
    uvicorn.run(create_app(), host="127.0.0.1", port=port, log_level="warning")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("mode", nargs="?", choices=["seed", "serve"], default="seed")
    ap.add_argument("--port", type=int, default=8123)
    ap.add_argument("--reseed", action="store_true", help="re-copy + re-seed before serving")
    args = ap.parse_args()
    if args.mode == "serve":
        serve(args.port, args.reseed)
    else:
        seed()


if __name__ == "__main__":
    main()
